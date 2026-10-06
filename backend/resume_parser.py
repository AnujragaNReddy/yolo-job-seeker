"""Turning an uploaded resume into a profile the matcher can use.

Text extraction for PDF, DOCX and plain text, then enough structure to search
with: skills, candidate job titles, years of experience and a seniority guess.

Kept to pypdf and python-docx rather than an NLP stack. A 512MB free instance
cannot hold spaCy or a transformer model alongside everything else, and for
this job the regex work is not meaningfully worse — the hard part is the
vocabulary in skills.py, not the parsing.

Nothing here phones home. The resume is read in-process, and what leaves the
machine is the search query built from it.
"""

import re
from pathlib import Path
from typing import Optional

import skills

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt", ".md"}

# Resumes are long; a cover letter prompt and the UI do not need all of it,
# and an enormous one is usually a portfolio export rather than a resume.
MAX_CHARS = 40_000

_EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+")
_PHONE = re.compile(r"(?:\+\d{1,3}[\s-]?)?(?:\(?\d{3,5}\)?[\s.-]?){2,4}\d{2,4}")
_LINKEDIN = re.compile(r"(?:linkedin\.com/in/)([A-Za-z0-9\-_%]+)", re.IGNORECASE)
_GITHUB = re.compile(r"(?:github\.com/)([A-Za-z0-9\-_]+)", re.IGNORECASE)

# "5+ years", "5-7 years", "over 5 years", "5 yrs"
_YEARS = re.compile(
    r"(\d{1,2})\s*(?:\+|plus)?\s*(?:-\s*\d{1,2}\s*)?(?:years?|yrs?)\b",
    re.IGNORECASE,
)

SENIORITY_ORDER = ["intern", "junior", "mid", "senior", "lead", "principal", "director"]

SENIORITY_HINTS = {
    "intern": ["intern", "internship", "trainee", "apprentice"],
    "junior": ["junior", "graduate", "entry level", "entry-level", "associate"],
    "senior": ["senior", "sr.", "sr "],
    "lead": ["lead", "team lead", "tech lead", "staff engineer"],
    "principal": ["principal", "architect", "distinguished"],
    "director": ["director", "head of", "vp ", "vice president", "cto", "chief"],
}

# Common role words, used to pull plausible titles out of the text. Matching a
# known vocabulary beats guessing from layout, which varies wildly between
# templates and is the main thing that makes resume parsing unreliable.
ROLE_WORDS = [
    "software engineer", "software developer", "full stack developer",
    "full-stack developer", "frontend developer", "front-end developer",
    "backend developer", "back-end developer", "web developer",
    "mobile developer", "android developer", "ios developer",
    "data engineer", "data scientist", "data analyst", "business analyst",
    "machine learning engineer", "ml engineer", "ai engineer",
    "devops engineer", "site reliability engineer", "cloud engineer",
    "platform engineer", "infrastructure engineer", "security engineer",
    "qa engineer", "test engineer", "automation engineer",
    "database administrator", "systems engineer", "network engineer",
    "product manager", "project manager", "program manager", "scrum master",
    "engineering manager", "technical lead", "solutions architect",
    "ux designer", "ui designer", "graphic designer", "product designer",
    "content writer", "technical writer", "digital marketer",
    "accountant", "financial analyst", "hr manager", "recruiter",
    "customer success manager", "support engineer", "sales executive",
    "consultant", "teacher", "operations manager",
]


class ResumeError(RuntimeError):
    """The file could not be read as a resume."""


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_SUFFIXES:
        raise ResumeError(
            f"{suffix or 'That file type'} is not supported. "
            f"Upload one of: {', '.join(sorted(SUPPORTED_SUFFIXES))}."
        )

    if suffix == ".pdf":
        text = _from_pdf(path)
    elif suffix == ".docx":
        text = _from_docx(path)
    else:
        text = path.read_text(encoding="utf-8", errors="replace")

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    if len(text) < 80:
        raise ResumeError(
            "Almost no text came out of that file. If it is a scanned or "
            "image-only PDF there is nothing to read — export a text PDF, or "
            "upload a .docx or .txt instead."
        )

    return text[:MAX_CHARS]


def _from_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ResumeError(
            "Reading PDFs needs pypdf. Run: pip install -r requirements.txt"
        ) from e

    try:
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as e:
        raise ResumeError(f"Could not read that PDF: {e}") from e


def _from_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as e:
        raise ResumeError(
            "Reading .docx needs python-docx. Run: pip install -r requirements.txt"
        ) from e

    try:
        document = docx.Document(str(path))
    except Exception as e:
        raise ResumeError(f"Could not read that .docx: {e}") from e

    parts = [p.text for p in document.paragraphs]

    # Plenty of resume templates put the entire history inside a table, which
    # document.paragraphs does not include at all.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)

    return "\n".join(part for part in parts if part and part.strip())


def guess_name(text: str) -> str:
    """The first line that looks like a person's name.

    A heuristic, and labelled as one everywhere it surfaces: it is used to
    address a cover letter, and the UI lets it be corrected.
    """
    for line in text.splitlines()[:8]:
        candidate = line.strip()

        if not candidate or len(candidate) > 45:
            continue
        if _EMAIL.search(candidate) or any(ch.isdigit() for ch in candidate):
            continue
        if any(word in candidate.lower() for word in
               ("resume", "curriculum", "vitae", "cv", "profile", "@", "http")):
            continue

        words = candidate.replace(",", " ").split()
        if 1 < len(words) <= 4 and all(w[:1].isalpha() for w in words):
            return " ".join(words).title()

    return ""


def guess_years(text: str) -> Optional[int]:
    """The largest credible "N years" in the document.

    The largest rather than the first: a resume mentions years against
    individual skills too, and the headline figure is the useful one. Capped
    at 45 because anything above that is a date range misread as a duration.
    """
    values = [int(match) for match in _YEARS.findall(text)]
    credible = [v for v in values if 0 < v <= 45]
    return max(credible) if credible else None


def guess_seniority(text: str, years: Optional[int]) -> str:
    head = text[:1500].lower()

    # Explicit wording wins over arithmetic — someone calling themselves a
    # lead is better evidence than a year count.
    for level in reversed(SENIORITY_ORDER):
        for hint in SENIORITY_HINTS.get(level, []):
            if hint in head:
                return level

    if years is None:
        return "mid"
    if years <= 1:
        return "junior"
    if years <= 5:
        return "mid"
    if years <= 9:
        return "senior"

    return "lead"


def guess_titles(text: str, limit: int = 6) -> list[str]:
    lowered = text.lower()
    found = []

    for role in ROLE_WORDS:
        if role in lowered:
            pretty = role.title().replace("Ui", "UI").replace("Ux", "UX")
            pretty = pretty.replace("Qa", "QA").replace("Ml", "ML").replace("Ai", "AI")
            pretty = pretty.replace("Hr", "HR")
            if pretty not in found:
                found.append(pretty)
        if len(found) >= limit:
            break

    return found


def build_profile(text: str, filename: str = "") -> dict:
    """Everything derived from the resume, in one record.

    Every field here is a guess from unstructured text, which is why the UI
    presents it as editable rather than as fact — a wrong skill list quietly
    ruins every match that follows, and the person reading it can fix in
    seconds what no parser will get right every time.
    """
    counts = skills.extract_counts(text)
    found = skills.extract(text)
    years = guess_years(text)

    def first(pattern: re.Pattern, group: int = 0) -> str:
        match = pattern.search(text)
        return match.group(group) if match else ""

    return {
        "filename": filename,
        "name": guess_name(text),
        "email": first(_EMAIL),
        "linkedin": first(_LINKEDIN, 1),
        "github": first(_GITHUB, 1),
        "skills": found,
        # Ordered by how often each appears, which is a decent proxy for what
        # the resume is actually about rather than what it merely mentions.
        "top_skills": [
            name for name, _ in sorted(counts.items(), key=lambda kv: -kv[1])
        ][:12],
        "titles": guess_titles(text),
        "years_experience": years,
        "seniority": guess_seniority(text, years),
        "characters": len(text),
    }
